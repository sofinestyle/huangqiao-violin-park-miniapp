from pathlib import Path
import base64, hashlib, json, shutil, sqlite3, subprocess, html
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
ASSETS = OUT / 'assets'
ASSETS.mkdir(exist_ok=True)
tracked = subprocess.check_output(['git','ls-files','-z'], cwd=ROOT).decode().split('\0')
protected = {p: hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in tracked if (ROOT/p).is_file() and not p.startswith(('docs/', 'AGENTS.md', 'README.md', 'skills/'))}
(OUT/'source-baseline.json').write_text(json.dumps({'gitHead':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),'protectedFiles':protected,'attachmentSha256':hashlib.sha256(Path('/Users/aaron/Desktop/P1-2.jpg').read_bytes()).hexdigest()},ensure_ascii=False,indent=2))

db=sqlite3.connect(f'file:{ROOT}/.local/huangqiao.sqlite?mode=ro',uri=True)
site=json.loads(db.execute("SELECT data FROM content WHERE kind='site' AND state='published' ORDER BY sort LIMIT 1").fetchone()[0])
media_id=site['images'][1].rsplit('/',1)[-1]
media_name=db.execute('SELECT stored_name FROM media WHERE id=?',(media_id,)).fetchone()[0]
db.close()
sources={'logo.png':ROOT/'images/yorray-logo.png','hero.png':ROOT/'.local/uploads'/media_name,'building.jpg':ROOT/'images/c1.jpg'}
manifest=[]
for dest,source in sources.items():
    shutil.copyfile(source,ASSETS/dest)
    manifest.append({'designFile':f'assets/{dest}','source':str(source),'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'usage':'原件副本，仅设计占位；权属及正式使用待审核' if dest!='logo.png' else '原标志副本，不重绘、不改色'})
(OUT/'素材来源.json').write_text(json.dumps({'assets':manifest,'companyIntro':site['intro'],'companyIntroSource':'本轮只读查询已发布站点资料','attachment':'/Users/aaron/Desktop/P1-2.jpg'},ensure_ascii=False,indent=2))
def data(name,mime):return 'data:'+mime+';base64,'+base64.b64encode((ASSETS/name).read_bytes()).decode()
logo=data('logo.png','image/png'); hero=data('hero.png','image/png'); building=data('building.jpg','image/jpeg')
company=html.escape(site['intro'])

page='''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>YorRay 品牌页设计稿</title><style>
:root{--ivory:#F7F3E9;--walnut:#3C2B24;--ink:#3C2C20;--secondary:#6C5948;--surface:#FFFDF7;--line:#E5DCCB;--scale:1;font-size:calc(16px * var(--scale))}*{box-sizing:border-box}html,body{margin:0;background:var(--ivory);color:var(--ink)}body{font-family:-apple-system,"PingFang SC","Microsoft YaHei",sans-serif}.serif,h1,h2,.glyph,.company{font-family:"Songti SC",STSong,"Noto Serif CJK SC",serif}button{font-family:inherit;cursor:pointer}button:focus-visible{outline:3px solid #80613B;outline-offset:4px}.page{width:100%;overflow-wrap:anywhere}.masthead{background:var(--walnut);color:var(--surface);text-align:center;padding:40px 24px 32px}.logo{display:block;width:176px;height:auto;margin:0 auto 16px}.brand-name{font-size:1rem;line-height:1.6;margin:0}.masthead h1{font-size:2.125rem;line-height:1.4;letter-spacing:.2em;font-weight:600;margin:16px 0 14px;padding-left:.2em}.masthead .tagline{font-size:.9375rem;line-height:1.8;margin:0;color:var(--surface)}.hero-photo{display:block;width:100%;height:240px;object-fit:cover;object-position:68% 56%}.section{padding:32px 24px}.eyebrow{font-size:.75rem;line-height:1.6;letter-spacing:.14em;color:var(--secondary);margin:0 0 12px}.section h2{font-size:1.5rem;font-weight:600;line-height:1.45;margin:0 0 20px}.section h2 span{display:block}.story p{font-size:1rem;line-height:1.85;margin:0 0 14px}.story p:last-child{margin:0}.meaning{padding-top:8px}.meaning h2{margin-bottom:12px}.meaning-row{display:flex;align-items:center;gap:24px;padding:22px 0}.meaning-row+.meaning-row{border-top:1px solid var(--line)}.glyph{font-size:3rem;font-weight:400;line-height:1.2;flex:0 0 60px}.meaning-copy{flex:1;min-width:0}.meaning-label{font-size:.8125rem;color:var(--secondary);margin:0 0 8px;line-height:1.6}.meaning-text{font-family:"Songti SC",STSong,serif;font-size:1.0625rem;line-height:1.8;margin:0}.poetry{padding-top:12px;padding-bottom:36px}.detail-photo{display:block;width:100%;aspect-ratio:4/3;object-fit:cover;object-position:74% 73%;border-radius:8px}.poem{margin:28px 0 0;text-align:center;font-size:1.375rem;line-height:1.85;font-weight:400}.poem p{margin:0}.building-photo{display:block;width:100%;aspect-ratio:3/2;object-fit:cover;object-position:50% 60%}.enterprise{padding-top:32px;padding-bottom:0}.enterprise h2{margin-bottom:20px}.company-label{margin:0 0 8px;font-size:.8125rem;line-height:1.6;color:var(--secondary)}.company{font-size:1.25rem;line-height:1.65;font-weight:600;margin:0 0 16px}.intro{font-size:.9375rem;line-height:1.9;margin:0;white-space:pre-line}.actions{display:flex;gap:12px;padding:28px 24px calc(32px + env(safe-area-inset-bottom,0px))}.action{flex:1;min-height:48px;border:0;border-radius:8px;font-size:.9375rem;line-height:1.6;padding:12px 8px;font-weight:500;background:#EDE5D7;color:var(--walnut)}.action.primary{background:var(--walnut);color:var(--surface)}.action:active{opacity:.75}.state-view{min-height:580px;padding:48px 24px;background:var(--ivory)}.state-view h1{font-size:1.5rem;line-height:1.5;margin:0 0 18px}.state-view p{font-size:1rem;line-height:1.85;margin:0 0 24px;color:var(--secondary)}.state-view button{padding:12px 22px;min-height:48px;background:var(--walnut);color:var(--surface);border:0;border-radius:8px;font-size:.9375rem}.failed-photo{display:flex;align-items:center;justify-content:center;text-align:center;background:#EDE5D7;color:var(--secondary);font-size:.875rem;line-height:1.8}.failed-photo.hero-photo{height:240px}.failed-photo.detail-photo{aspect-ratio:4/3}.failed-photo.building-photo{aspect-ratio:3/2}.sr-feedback{position:absolute;width:1px;height:1px;overflow:hidden;clip-path:inset(50%)}[hidden]{display:none!important}@media(max-width:340px){.section{padding-left:20px;padding-right:20px}.masthead{padding-left:20px;padding-right:20px}.masthead h1{font-size:1.875rem}.actions{padding-left:20px;padding-right:20px;flex-direction:column}.meaning-row{gap:16px}.meaning h2{font-size:1.375rem}}html[data-large="true"] .actions{flex-direction:column}html[data-large="true"] .masthead{padding-top:40px}
</style></head><body><main class="page" id="page">
<header class="masthead" id="brand"><img class="logo" src="__LOGO__" alt="YorRay"><p class="brand-name">瑜瑞 YorRay</p><h1>以诗入琴</h1><p class="tagline">名字是世界上最短的诗。</p></header>
<img class="hero-photo" src="__HERO__" alt="提琴摄影设计占位" id="hero-photo">
<section class="section story" id="story"><p class="eyebrow">BRAND CULTURE</p><h2><span>名字，让想象</span><span>有了形状</span></h2><p>名字是世界上最短的诗。</p><p>当人们开始为万物命名，日月星辰，山川湖海，从此世界有了模样、想象有了形状。</p></section>
<section class="section meaning" id="meaning"><h2>瑜之光彩，瑞之信意</h2><div class="meaning-row"><span class="glyph">瑜</span><div class="meaning-copy"><p class="meaning-label">光彩</p><p class="meaning-text">瑜，玉之光彩也。</p></div></div><div class="meaning-row"><span class="glyph">瑞</span><div class="meaning-copy"><p class="meaning-label">信意</p><p class="meaning-text">瑞，以玉为信也。</p></div></div></section>
<section class="section poetry" id="poetry"><img class="detail-photo" src="__HERO__" alt="提琴细节设计占位"><div class="poem serif"><p>以诗入琴，</p><p>万物有了雅称，</p><p>世界便有了美的意义。</p></div></section>
<img class="building-photo" src="__BUILDING__" alt="项目现有城市客厅外景设计占位" id="building-photo">
<section class="section enterprise" id="enterprise"><p class="eyebrow">BRAND &amp; COMPANY</p><h2>品牌与企业</h2><p class="company-label">运营主体</p><p class="company">江苏黄桥乐器文化产业园<br>投资发展有限公司</p><p class="intro">__COMPANY__</p></section>
<footer class="actions" id="actions"><button class="action primary" type="button" data-action="instruments">查看乐器</button><button class="action" type="button" data-action="consult">联系咨询</button></footer>
</main><main class="state-view" id="state-view" hidden></main><p class="sr-feedback" id="feedback" aria-live="polite"></p>
<script>
const qs=new URLSearchParams(location.search),state=qs.get('state')||'normal',large=qs.get('large')==='1';if(large){document.documentElement.style.setProperty('--scale','1.2');document.documentElement.dataset.large='true'}
const stateView=document.querySelector('#state-view');if(['loading','error','empty'].includes(state)){document.querySelector('#page').hidden=true;stateView.hidden=false;const text={loading:['品牌资料加载中…','正在读取品牌资料，请稍候。'],error:['品牌资料暂未加载','加载失败，请稍后重试。'],empty:['品牌资料暂未发布','品牌资料暂未发布，请稍后再试。']}[state];const h=document.createElement('h1');h.textContent=text[0];stateView.append(h);const p=document.createElement('p');p.textContent=text[1];stateView.append(p);if(state==='error'){const b=document.createElement('button');b.textContent='重新加载';b.onclick=()=>{if(parent!==window)parent.postMessage({kind:'retry'},'*');else{qs.delete('state');location.search=qs.toString()}};stateView.append(b)}}if(state==='image-error')for(const img of document.querySelectorAll('img:not(.logo)')){const box=document.createElement('div');box.className=img.className+' failed-photo';box.textContent='图片暂未加载';img.replaceWith(box)}
for(const button of document.querySelectorAll('[data-action]'))button.onclick=()=>{const action=button.dataset.action;const message=action==='instruments'?'入口示意：正式页面将进入乐器栏目。本稿未执行跳转。':'入口示意：正式页面将进入现有咨询页。本稿不提交咨询。';document.querySelector('#feedback').textContent=message;if(parent!==window)parent.postMessage({kind:'action',action,message},'*');else alert(message)};
</script></body></html>'''.replace('__LOGO__',logo).replace('__HERO__',hero).replace('__BUILDING__',building).replace('__COMPANY__',company)
page=page.replace('<img class="detail-photo"','<div class="detail-frame"><img class="detail-photo"').replace('alt="提琴细节设计占位">','alt="提琴细节设计占位"></div>')
page=page.replace('</style>','.detail-frame{position:relative;width:100%;aspect-ratio:4/3;overflow:hidden;border-radius:8px}.detail-frame .detail-photo{position:absolute;inset:0;width:100%;height:100%;aspect-ratio:auto;border-radius:0}.detail-frame img.detail-photo{transform:scale(1.7);transform-origin:76% 67%}</style>')
(OUT/'page.html').write_text(page,encoding='utf8')

review='''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>YorRay 品牌页面设计审阅</title><style>
*{box-sizing:border-box}body{margin:0;color:#302820;background:#EFECE6;font-family:-apple-system,"PingFang SC","Microsoft YaHei",sans-serif}header{padding:24px 32px 20px;border-bottom:1px solid #D8D1C7;background:#F8F6F1}header p{margin:0;font-size:13px;color:#615548;line-height:1.8}h1{font-size:24px;font-weight:600;line-height:1.5;margin:0 0 6px}.workspace{display:grid;grid-template-columns:250px minmax(340px,1fr) 280px;max-width:1300px;margin:0 auto;gap:32px;padding:28px 32px}.controls,.notes{line-height:1.8;font-size:14px}h2{font-size:16px;margin:0 0 12px}h3{font-size:14px;margin:20px 0 8px}.group{margin-bottom:24px}.choices{display:flex;flex-wrap:wrap;gap:8px}button{min-height:40px;padding:8px 12px;border:1px solid #CFC7BC;background:#F8F6F1;color:#3C2B24;font:inherit;font-size:13px;border-radius:5px;cursor:pointer}button[aria-pressed="true"]{background:#3C2B24;color:#FFFDF7;border-color:#3C2B24}button:focus-visible,a:focus-visible{outline:3px solid #80613B;outline-offset:3px}select{font:inherit;width:100%;min-height:44px;padding:8px;background:#F8F6F1;border:1px solid #CFC7BC;border-radius:5px}label{display:flex;gap:9px;align-items:center;min-height:44px}input{width:18px;height:18px}.preview{display:flex;align-items:center;flex-direction:column;min-width:0}.preview-label{font-size:13px;color:#615548;margin:0 0 12px}iframe{border:0;background:#F7F3E9;box-shadow:0 8px 28px #35271B19;height:calc(100vh - 190px);min-height:540px;max-height:860px;max-width:100%}.notes p{margin:0 0 14px}.notes ul{padding-left:18px;margin:0 0 16px}.notes li{margin:0 0 9px}a{color:#56432F;text-underline-offset:3px}.feedback{border-top:1px solid #D8D1C7;padding-top:16px;margin-top:24px!important;color:#3C2B24}.swatches{display:flex;gap:10px}.swatch{width:28px;height:28px;border-radius:50%;border:1px solid #d7cfc5}.caption{font-size:12px;color:#615548}.note-footer{font-size:12px;line-height:1.8;color:#615548;margin-top:24px} @media(max-width:1100px){.workspace{grid-template-columns:220px minmax(320px,1fr);gap:24px}.notes{grid-column:1/-1;display:grid;grid-template-columns:1fr 1fr;gap:24px}}@media(max-width:660px){header{padding:20px}.workspace{display:flex;flex-direction:column;padding:20px 12px;gap:24px}.controls{padding:0 8px}.controls .group{margin-bottom:16px}.notes{display:block;padding:0 8px}iframe{height:760px;min-height:540px}.preview{width:100%}.choices button{min-height:44px}h1{font-size:21px}}@media print{header,.controls,.notes{display:none}.workspace{display:block;padding:0}.preview-label{display:none}iframe{height:2600px;box-shadow:none}}
</style></head><body><header><h1>YorRay 品牌页面设计审阅</h1><p>V1.0 · 2026年10月6日 · 主标题已确认「以诗入琴」 · 仅设计稿，未修改小程序</p></header>
<main class="workspace"><aside class="controls"><div class="group"><h2>页面宽度</h2><div class="choices" id="width-controls"><button type="button" data-width="320" aria-pressed="false">320</button><button type="button" data-width="375" aria-pressed="false">375</button><button type="button" data-width="390" aria-pressed="true">390</button><button type="button" data-width="430" aria-pressed="false">430</button></div><p class="caption">单位：逻辑像素。窄屏自动换行。</p></div><div class="group"><h2>页面状态</h2><select id="state" aria-label="页面状态"><option value="normal">正常页面</option><option value="loading">加载中</option><option value="error">读取失败</option><option value="empty">资料未发布</option><option value="image-error">图片未加载</option></select></div><div class="group"><label><input id="large" type="checkbox">正文放大至120%</label></div><div class="group"><h2>章节定位</h2><div class="choices"><button data-section="brand" type="button">品牌首屏</button><button data-section="story" type="button">名字的诗意</button><button data-section="meaning" type="button">瑜与瑞</button><button data-section="enterprise" type="button">品牌与企业</button></div></div><p class="note-footer">页面可在中间窗口滚动。所有状态切换和按钮示意只作用于本设计稿。</p><p><a href="YorRay品牌页面最终设计方案_V1.0.docx">下载详细Word方案</a></p><p><a href="品牌页完整设计稿_390.png" target="_blank">查看完整长图</a></p></aside>
<section class="preview" aria-label="品牌页设计预览"><p class="preview-label" id="preview-label">390px · 正常页面</p><iframe id="frame" width="390" title="YorRay 品牌页设计稿"></iframe></section>
<aside class="notes"><div><h2>设计说明</h2><p>以品牌文化为主线，完整保留企业背景与两个现有服务入口。区分「黄桥琴旅」小程序名称、「瑜瑞 YorRay」品牌及运营主体。</p><h3>阅读顺序</h3><ol><li>品牌识别与主标题</li><li>名字的诗意</li><li>瑜与瑞的寓意</li><li>以诗入琴的品牌表达</li><li>完整企业介绍与服务入口</li></ol><h3>视觉语言</h3><div class="swatches" aria-label="暖米白、深胡桃木、装饰金"><span class="swatch" style="background:#F7F3E9"></span><span class="swatch" style="background:#3C2B24"></span><span class="swatch" style="background:#B99A68"></span></div><p class="caption">中文衬线标题 · 清楚正文 · 克制留白</p></div><div><h2>素材与范围</h2><p>标志沿用项目原件；提琴图沿用当前站点图片；外景取自项目现有城市客厅图片。照片为设计占位，权属与品牌代表性尚未完成正式审核。</p><p>附件中的品牌释义作为品牌赋义呈现。未添加销量、认证、工厂归属或品质承诺。</p><p>本稿未连接业务接口，未修改小程序、后台或数据库。未执行微信原生及真机验收。</p><p class="feedback" id="feedback" aria-live="polite">点击页尾按钮，可查看已有入口的用途说明。</p></div></aside></main>
<script>
const frame=document.querySelector('#frame'),state=document.querySelector('#state'),large=document.querySelector('#large'),preview=document.querySelector('#preview-label');const base=__PAGE_JSON__;let width=390;
function render(){const params=new URLSearchParams({state:state.value,large:large.checked?'1':'0'});frame.style.width=width+'px';frame.srcdoc=base.replace("const qs=new URLSearchParams(location.search)","const qs=new URLSearchParams("+JSON.stringify(params.toString())+")");preview.textContent=width+'px · '+state.options[state.selectedIndex].text+(large.checked?' · 120%文字':'');for(const b of document.querySelectorAll('[data-width]'))b.setAttribute('aria-pressed',String(Number(b.dataset.width)===width))}document.querySelectorAll('[data-width]').forEach(b=>b.onclick=()=>{width=Number(b.dataset.width);render()});state.onchange=render;large.onchange=render;document.querySelectorAll('[data-section]').forEach(b=>b.onclick=()=>{if(state.value!=='normal'&&state.value!=='image-error'){state.value='normal';render();frame.onload=()=>{frame.contentDocument.getElementById(b.dataset.section)?.scrollIntoView({behavior:'smooth'});frame.onload=null}}else frame.contentDocument.getElementById(b.dataset.section)?.scrollIntoView({behavior:'smooth'})});window.addEventListener('message',e=>{if(e.source!==frame.contentWindow)return;if(e.data.kind==='action')document.querySelector('#feedback').textContent=e.data.message;if(e.data.kind==='retry'){state.value='normal';render();document.querySelector('#feedback').textContent='重试示意：已返回正常设计稿，未请求服务器。'}});render();
</script></body></html>'''.replace('__PAGE_JSON__',json.dumps(page,ensure_ascii=False).replace('</script>',r'<\/script>'))
(OUT/'品牌页面设计审阅.html').write_text(review,encoding='utf8')

def font(run,name='Songti SC',size=11,bold=False):
    run.font.name=name;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    rpr=run._element.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rpr.append(rf)
    for k in ['ascii','hAnsi','eastAsia']:rf.set(qn('w:'+k),name)

doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8);sec.header_distance=Inches(.28);sec.footer_distance=Inches(.3)
for grid in sec._sectPr.xpath('./w:docGrid'):sec._sectPr.remove(grid)
for style in doc.styles:
    for border in style._element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for sn in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[sn];st.font.name='PingFang SC' if sn!='Normal' else 'Songti SC';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt({'Normal':11,'Title':23,'Heading 1':17,'Heading 2':12.5,'Heading 3':12}[sn]);st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(8);st.paragraph_format.line_spacing=1.25
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),st.font.name)
    grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');st._element.get_or_add_pPr().append(grid)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
for sn,leading in [('Normal',17),('Title',32),('Heading 1',25),('Heading 2',19),('Heading 3',18)]:doc.styles[sn].paragraph_format.line_spacing=Pt(leading)
doc.styles['Title'].paragraph_format.space_after=Pt(14)
for sn in ['Heading 1','Heading 2','Heading 3']:doc.styles[sn].paragraph_format.keep_with_next=True;doc.styles[sn].paragraph_format.space_before=Pt(12)
header=sec.header.paragraphs[0];font(header.add_run('YorRay 品牌页面设计方案'),name='PingFang SC',size=9)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(footer.add_run('V1.0  |  2026年10月6日  |  '),name='PingFang SC',size=9);f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.title='YorRay品牌页面最终设计方案';doc.core_properties.subject='仅详细设计方案和页面设计稿 不修改小程序';doc.core_properties.author='';doc.core_properties.last_modified_by=''
lines=(OUT/'YorRay品牌页面最终设计方案_V1.0.md').read_text().splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line=='<!-- PAGE -->':i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(set(v)<=set('-: ') for v in row):rows.append(row)
            i+=1
        table=doc.add_table(rows=0, cols=len(rows[0]));table.autofit=False
        widths=([1.05,2.65,3.2] if len(rows[0])==3 else [2.0,4.9]);total=sum(widths);widths=[v/total*6.9 for v in widths]
        for j,width in enumerate(widths):table.columns[j].width=Inches(width)
        for k,row in enumerate(rows):
            cells=table.add_row().cells
            for j,value in enumerate(row):
                cells[j].width=Inches(widths[j]);cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p=cells[j].paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=Pt(14.5);r=p.add_run(value);font(r,name='PingFang SC',size=9.5,bold=(k==0))
                tcPr=cells[j]._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'404040' if k==0 else ('F2F2F2' if k%2==0 else 'FFFFFF'));tcPr.append(shd)
                if k==0:r.font.color.rgb=RGBColor(255,255,255)
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                tcPr.append(borders);margins=OxmlElement('w:tcMar')
                for edge,val in [('top','75'),('bottom','75'),('left','100'),('right','100')]:
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');margins.append(el)
                tcPr.append(margins)
            if k==0:
                trPr=table.rows[k]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');trPr.append(rep)
            trPr=table.rows[k]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
        doc.add_paragraph().paragraph_format.space_after=Pt(1)
        continue
    if line.startswith('# '):p=doc.add_paragraph(line[2:],style='Title')
    elif line.startswith('## '):p=doc.add_paragraph(line[3:],style='Heading 1')
    elif line.startswith('### '):p=doc.add_paragraph(line[4:],style='Heading 2')
    else:
        p=doc.add_paragraph(line);p.paragraph_format.keep_together=True
        later=next((s.strip() for s in lines[i+1:] if s.strip()),'')
        if len(line)<45 and later and not later.startswith(('#','|','<!--')):p.paragraph_format.keep_with_next=True
    i+=1
doc.save(OUT/'YorRay品牌页面最终设计方案_V1.0.docx')
print(json.dumps({'directory':str(OUT),'companyIntroCharacters':len(site['intro']),'protectedFiles':len(protected),'designHtmlBytes':len(review.encode()),'docxCreated':True},ensure_ascii=False))
